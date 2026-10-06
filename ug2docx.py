#!/usr/bin/env python3
"""
ug2docx.py - turn an Ultimate Guitar chord sheet into a clean Word document,
and transpose chords in those documents (or in your existing ones) afterwards.

Setup (once):
    pip install python-docx pillow requests reportlab
    (optional, helps if Ultimate Guitar blocks plain requests:  pip install cloudscraper)
    pip install reportlab                              # needed for --pdf

Make a document:
    python ug2docx.py make "https://tabs.ultimate-guitar.com/tab/....-chords-1234567"
    python ug2docx.py make URL --transpose -2          # 2 semitones down
    python ug2docx.py make URL --flats                 # use Bb instead of A#
    python ug2docx.py make URL --pdf                   # also save a PDF
    python ug2docx.py make URL --size 12               # fixed size (default: auto-fit 8-12 pt)
    python ug2docx.py make URL --keep-capo             # keep the guitar/capo chords
    python ug2docx.py make --html saved_page.html      # if the download is blocked,
                                                       # save the page in your browser
                                                       # (Ctrl/Cmd+S) and use this
    python ug2docx.py make --text pasted.txt --title "Song" --artist "Band"

Transpose a document you already have (also works on your older ones):
    python ug2docx.py transpose "Song - Band.docx" +2
    python ug2docx.py transpose "Song - Band.docx" -3 --flats -o lower.docx
    python ug2docx.py transpose "Song - Band.docx" +2 --pdf        # also save a PDF

How the chords stay above the right words
    Chords are placed with Word tab stops at the exact horizontal position of
    the syllable they belong to (measured with real Arial metrics), instead of
    padding with spaces. Alignment therefore does not depend on spaces, and
    transposing never shifts the chords sideways.

How sections stay together
    Every line of a section (Verse, Chorus, ...) is set to "keep with next",
    so if the section does not fit on the page Word moves all of it to the
    next page. This keeps working after you edit the document.
"""

import argparse
import html as htmlmod
import json
import os
import re
import sys

# --------------------------------------------------------------------------
# Layout settings (edit to taste)
# --------------------------------------------------------------------------
FONT_NAME = "Arial"
BODY_SIZE = 12
TITLE_SIZE = 18
LINE_SPACING = 1.15
SPACE_AFTER_PT = 6
PAGE_WIDTH_CM, PAGE_HEIGHT_CM = 21.0, 29.7      # A4 (US Letter: 21.59, 27.94)
MARGIN_CM = 2.54                                 # 1 inch
MIN_CHORD_GAP_PT = 4                             # minimum gap between two chords
MIN_SIZE, MAX_SIZE, SIZE_STEP = 8, 12, 0.5       # auto-fit tries these body sizes

# --------------------------------------------------------------------------
# Chord parsing / transposition
# --------------------------------------------------------------------------
SHARP = ["C", "C#", "D", "D#", "E", "F", "F#", "G", "G#", "A", "A#", "B"]
FLAT = ["C", "Db", "D", "Eb", "E", "F", "Gb", "G", "Ab", "A", "Bb", "B"]
NOTE_INDEX = {n: i for i, n in enumerate(SHARP)}
NOTE_INDEX.update({n: i for i, n in enumerate(FLAT)})
NOTE_INDEX.update({"Cb": 11, "Fb": 4, "E#": 5, "B#": 0})

CHORD_RE = re.compile(
    r"^([A-G][#b]?)"                                              # root
    r"((?:m|min|maj|dim|aug|sus|add|M|\+|°|ø|-|\d|\(|\)|#|b|,)*)"  # quality/extensions
    r"(?:/([A-G][#b]?))?$"                                        # optional bass note
)


def is_chord(token):
    return bool(CHORD_RE.match(token))


def transpose_chord(chord, steps, flats):
    m = CHORD_RE.match(chord)
    if not m:
        return chord
    root, suffix, bass = m.groups()
    names = FLAT if flats else SHARP
    out = names[(NOTE_INDEX[root] + steps) % 12] + suffix
    if bass:
        out += "/" + names[(NOTE_INDEX[bass] + steps) % 12]
    return out


def prefers_flats(chords):
    """Guess from the chords: flats if there are flats and no sharps."""
    has_flat = any(re.match(r"^[A-G]b", c) or re.search(r"/[A-G]b", c) for c in chords)
    has_sharp = any(re.match(r"^[A-G]#", c) or re.search(r"/[A-G]#", c) for c in chords)
    return has_flat and not has_sharp


def auto_flats(chords):
    """Pick flat or sharp spelling by guessing the key the chords fit best.
    Flats only for F, Bb, Eb, Ab major (and their relative minors); everything
    else, including Db/C# and Gb/F#, stays in sharps like Ultimate Guitar."""
    roots = [NOTE_INDEX[CHORD_RE.match(c).group(1)] for c in chords if is_chord(c)]
    if not roots:
        return False
    best = max(
        (sum(r in {(t + i) % 12 for i in (0, 2, 4, 5, 7, 9, 11)} for r in roots)
         + (0.5 if roots[0] in (t, (t + 9) % 12) else 0), t)
        for t in range(12))
    return best[1] in (5, 10, 3, 8)


# --------------------------------------------------------------------------
# Fetching from Ultimate Guitar
# --------------------------------------------------------------------------
BROWSER_HEADERS = {
    "User-Agent": ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}


def fetch_html(url):
    try:
        import cloudscraper  # type: ignore
        session = cloudscraper.create_scraper()
    except ImportError:
        import requests
        session = requests.Session()
    resp = session.get(url, headers=BROWSER_HEADERS, timeout=30)
    if resp.status_code != 200:
        raise RuntimeError(
            f"Ultimate Guitar answered HTTP {resp.status_code}. "
            "Open the page in your browser, save it (Ctrl/Cmd+S) and run:\n"
            "    python ug2docx.py make --html saved_page.html")
    return resp.text


def _find_key(obj, key):
    if isinstance(obj, dict):
        if key in obj:
            return obj[key]
        for v in obj.values():
            r = _find_key(v, key)
            if r is not None:
                return r
    elif isinstance(obj, list):
        for v in obj:
            r = _find_key(v, key)
            if r is not None:
                return r
    return None


def parse_ug_html(page):
    m = re.search(r'class="js-store"[^>]*data-content="([^"]*)"', page)
    if not m:
        m = re.search(r'data-content="([^"]*)"[^>]*class="js-store"', page)
    if not m:
        raise RuntimeError("Could not find the chord data in the page "
                           "(Ultimate Guitar may have changed its layout, or you got a block page).")
    data = json.loads(htmlmod.unescape(m.group(1)))
    content = _find_key(data, "wiki_tab")
    content = content.get("content") if isinstance(content, dict) else None
    if not content:
        raise RuntimeError("No chord content found on that page.")
    title = _find_key(data, "song_name") or "Unknown Song"
    artist = _find_key(data, "artist_name") or "Unknown Artist"
    capo = _find_key(_find_key(data, "tab_view") or {}, "capo")
    try:
        capo = int(capo or 0)
    except (TypeError, ValueError):
        capo = 0
    return title, artist, content, capo


# --------------------------------------------------------------------------
# Parsing the chord sheet text
# --------------------------------------------------------------------------
TAG_RE = re.compile(r"\[/?(?:tab|ch)\]")
HEADER_RE = re.compile(r"^\s*\[(?!/?(?:ch|tab)\])([^\]]+)\]\s*$")
CH_SPLIT = re.compile(r"\[ch\](.*?)\[/ch\]")


def parse_chord_line(line):
    """Return list of (column, chord) for a line containing [ch] tags, or
    None if the line also contains lyrics/other text."""
    plain, chords, pos = "", [], 0
    for m in CH_SPLIT.finditer(line):
        plain += line[pos:m.start()]
        chords.append((len(plain), m.group(1).strip()))
        plain += m.group(1)
        pos = m.end()
    plain += line[pos:]
    if not chords:
        return None
    # a "chord line" has nothing but chords and whitespace
    check = plain
    for _, c in chords:
        check = check.replace(c, "", 1)
    if check.strip(" \t|-()x0123456789"):
        return None
    return chords


def split_preface(content):
    """Text before the first [Section] header -> (preface, body)."""
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    lines = content.split("\n")
    for i, l in enumerate(lines):
        if HEADER_RE.match(l):
            return "\n".join(lines[:i]), "\n".join(lines[i:])
    return "", content


def detect_capo(content, meta_capo=0):
    """Capo fret from the page data, else from a 'Capo: 3rd fret' line."""
    if meta_capo:
        return meta_capo
    preface, _ = split_preface(content)
    m = re.search(r"capo\D{0,15}?(\d{1,2})", TAG_RE.sub("", preface), re.I)
    return int(m.group(1)) if m else 0


def build_sections(content):
    """-> list of (header or None, [items]) where items are
         ("chords", [(col, chord)...])  chord-only line
         ("pair", [(col, chord)...], lyric)
         ("lyric", text)"""
    _, content = split_preface(content)          # drop everything before the first [Section]
    content = re.sub(r"\[/?tab\]", "", content)
    lines = content.split("\n")
    has_headers = any(HEADER_RE.match(l) for l in lines)

    sections, cur_header, cur_items = [], None, []

    def flush():
        nonlocal cur_header, cur_items
        if cur_items or cur_header:
            sections.append((cur_header, cur_items))
        cur_header, cur_items = None, []

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        hm = HEADER_RE.match(line)
        if hm:
            flush()
            cur_header = hm.group(1).strip()
            cur_header = f"[{cur_header}]"
            i += 1
            continue
        if not line.strip():
            if not has_headers:
                flush()
            i += 1
            continue
        chords = parse_chord_line(line)
        if chords is not None:
            # lyric = next non-empty line that isn't a chord line / header
            nxt = lines[i + 1].rstrip() if i + 1 < len(lines) else ""
            if (nxt.strip() and not HEADER_RE.match(nxt)
                    and parse_chord_line(nxt) is None and "[ch]" not in nxt):
                cur_items.append(("pair", chords, nxt))
                i += 2
                continue
            cur_items.append(("chords", chords))
            i += 1
            continue
        cur_items.append(("lyric", TAG_RE.sub("", line)))
        i += 1
    flush()
    return sections


# --------------------------------------------------------------------------
# Text measurement (real Arial widths so chords land above the right letter)
# --------------------------------------------------------------------------
_FONT_CANDIDATES = {
    False: [
        "/Library/Fonts/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/liberation/LiberationSans-Regular.ttf",
    ],
    True: [
        "/Library/Fonts/Arial Bold.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        "/usr/share/fonts/liberation/LiberationSans-Bold.ttf",
    ],
}


class Measurer:
    SCALE = 1000

    def __init__(self):
        self.fonts = {}
        self.rl = None
        try:
            from reportlab.pdfbase.pdfmetrics import stringWidth
            self.rl = stringWidth        # Helvetica has exactly Arial's letter widths
            return
        except ImportError:
            pass
        try:
            from PIL import ImageFont
            for bold, paths in _FONT_CANDIDATES.items():
                for p in paths:
                    if os.path.exists(p):
                        self.fonts[bold] = ImageFont.truetype(p, self.SCALE)
                        break
        except ImportError:
            pass
        if len(self.fonts) < 2:
            print("WARNING: no Arial (or Liberation Sans) font file found - chord "
                  "positions will be approximate.", file=sys.stderr)

    def width(self, text, bold=False, size=BODY_SIZE):
        """Width in points."""
        if self.rl:
            return self.rl(text, "Helvetica-Bold" if bold else "Helvetica", size)
        font = self.fonts.get(bold)
        if font is None:
            return len(text) * size * (0.58 if bold else 0.53)
        return font.getlength(text) * size / self.SCALE


# --------------------------------------------------------------------------
# Word document creation
# --------------------------------------------------------------------------
def prepare(sections, steps=0, flats=None):
    """Transpose chord names and normalise lyric lines. Returns sections of
         ("pair", [(col, chord)], lyric) | ("chords", [chord]) | ("lyric", text)"""
    allc = [c for _, items in sections for it in items
            if it[0] in ("chords", "pair") for _, c in it[1]]
    if flats is None and steps:
        flats = auto_flats([transpose_chord(c, steps, False) for c in allc if is_chord(c)])
    rename = bool(steps) or flats is not None

    def tc(c):
        return transpose_chord(c, steps, bool(flats)) if rename and is_chord(c) else c

    out = []
    for header, items in sections:
        new = []
        for it in items:
            if it[0] == "pair":
                raw = it[2].replace("\t", "    ").rstrip()
                lead = len(raw) - len(raw.lstrip())
                new.append(("pair", [(max(0, c - lead), tc(ch)) for c, ch in it[1]], raw.lstrip()))
            elif it[0] == "chords":
                new.append(("chords", [tc(c) for _, c in it[1]]))
            else:
                new.append(it)
        out.append((header, new))
    return out


def place_chords(chords, lyric, meas, max_x, size=BODY_SIZE):
    """[(col, chord)] -> [(x_in_points, chord)] so chords sit over their letters."""
    out, prev_end = [], None
    for col, chord in chords:
        padded = lyric if col <= len(lyric) else lyric + " " * (col - len(lyric))
        x = meas.width(padded[:col], size=size)
        if prev_end is not None:
            x = max(x, prev_end + MIN_CHORD_GAP_PT)
        x = min(x, max_x - 10)
        out.append((x, chord))
        prev_end = x + meas.width(chord, bold=True, size=size)
    return out


def make_docx(title, artist, sections, path, transpose=0, flats=None, size=BODY_SIZE):
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    sections = prepare(sections, transpose, flats)
    meas = Measurer()
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(PAGE_WIDTH_CM), Cm(PAGE_HEIGHT_CM)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(MARGIN_CM)
    text_width_pt = Cm(PAGE_WIDTH_CM - 2 * MARGIN_CM).pt

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(size)
    rfonts = normal.element.get_or_add_rPr().find(qn("w:rFonts"))
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT_NAME)
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.space_after = Pt(SPACE_AFTER_PT)
    normal.paragraph_format.space_before = Pt(0)

    def run(par, text, bold=False, sz=None):
        r = par.add_run(text)
        r.font.name = FONT_NAME
        r.font.size = Pt(sz or size)
        r.bold = bold
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        return r

    # Title: "Song Title - Artist", Arial 18 bold (always 18, whatever the body size)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run(p, f"{title} \u2013 {artist}", bold=True, sz=TITLE_SIZE)

    def chord_par(chords, lyric):
        par = doc.add_paragraph()
        for x, chord in place_chords(chords, lyric, meas, text_width_pt, size):
            if x < 1:                       # chord at the very start of the line
                run(par, chord, bold=True)
            else:
                par.paragraph_format.tab_stops.add_tab_stop(Pt(round(x, 1)), WD_TAB_ALIGNMENT.LEFT)
                run(par, "\t" + chord, bold=True)
        return par

    for n, (header, items) in enumerate(sections):
        block = []
        if header:
            hp = doc.add_paragraph()
            run(hp, header)
            block.append(hp)
        for it in items:
            if it[0] == "pair":
                block.append(chord_par(it[1], it[2]))
                lp = doc.add_paragraph()
                run(lp, it[2])
                block.append(lp)
            elif it[0] == "chords":
                cp = doc.add_paragraph()
                run(cp, " ".join(it[1]), bold=True)
                block.append(cp)
            else:
                lp = doc.add_paragraph()
                run(lp, it[1])
                block.append(lp)
        # keep the whole section on one page where possible
        for idx, par in enumerate(block):
            par.paragraph_format.keep_together = True
            par.paragraph_format.keep_with_next = idx < len(block) - 1
        if n < len(sections) - 1:
            doc.add_paragraph()              # empty spacer between sections

    doc.save(path)
    return path


PT_PER_CM = 72 / 2.54
_line = lambda size: size * 1.149 * LINE_SPACING            # Arial line height x spacing


def _default_geom():
    m = MARGIN_CM * PT_PER_CM
    return (PAGE_WIDTH_CM * PT_PER_CM, PAGE_HEIGHT_CM * PT_PER_CM, m, m, m)


def _layout(sections, body_size, has_title, geom):
    """Decide which page and height every paragraph lands on.
    sections: list of (units, gap_rows); unit = list of paragraphs kept together.
    Returns (placements [(page, y_top, paragraph)], page_count)."""
    W, H, left, top, bottom = geom
    row = _line(body_size) + SPACE_AFTER_PT
    page, y = 0, H - top - (_line(TITLE_SIZE) if has_title else 0)
    placed = []
    for units, gap in sections:
        total = sum(len(u) for u in units) * row - SPACE_AFTER_PT
        if units and total > y - bottom and total <= H - top - bottom:
            page, y = page + 1, H - top            # whole section moves to the next page
        for u in units:
            need = len(u) * row - SPACE_AFTER_PT
            if y - need < bottom:
                page, y = page + 1, H - top
            for par in u:
                placed.append((page, y, par))
                y -= row
        y -= row * gap
    return placed, page + 1


def _write_pdf(path, doc_title, title_text, sections, geom=None, body_size=BODY_SIZE):
    """Low-level PDF writer shared by make_pdf and docx_to_pdf.
    paragraph = list of (x_pt, text, bold, size). geom = (W, H, left, top, bottom) in pt."""
    from reportlab.pdfgen import canvas

    geom = geom or _default_geom()
    W, H, left, top, bottom = geom
    clean = lambda t: t.encode("cp1252", "replace").decode("cp1252")
    placed, pages = _layout(sections, body_size, bool(title_text), geom)

    c = canvas.Canvas(path, pagesize=(W, H))
    c.setTitle(doc_title)

    def draw(x, y_top, text, bold, size):
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.drawString(left + x, y_top - _line(size) + 0.27 * size, clean(text))

    if title_text:
        draw(0, H - top, title_text, True, TITLE_SIZE)
    page = 0
    for pg, y, par in placed:
        while page < pg:
            c.showPage()
            page += 1
        for x, text, bold, size in par:
            draw(x, y, text, bold, size)
    c.save()
    return path


def _pdf_sections(prepared, size, meas, text_w):
    """Prepared sections -> drawing paragraphs at a given body size.
    Also returns the widest line (points), to spot lines that would wrap."""
    out, widest = [], 0.0
    for header, items in prepared:
        units = []
        if header:
            units.append([[(0, header, False, size)]])
            widest = max(widest, meas.width(header, False, size))
        for it in items:
            if it[0] == "pair":
                placed = place_chords(it[1], it[2], meas, text_w, size)
                units.append([[(x, ch, True, size) for x, ch in placed], [(0, it[2], False, size)]])
                widest = max(widest, meas.width(it[2], False, size))
                if placed:
                    widest = max(widest, placed[-1][0] + meas.width(placed[-1][1], True, size))
            elif it[0] == "chords":
                t = " ".join(it[1])
                units.append([[(0, t, True, size)]])
                widest = max(widest, meas.width(t, True, size))
            else:
                units.append([[(0, it[1], False, size)]])
                widest = max(widest, meas.width(it[1], False, size))
        out.append((units, 1))
    return out, widest


def plan_layout(title, artist, sections, transpose=0, flats=None, autofit=True, size=BODY_SIZE):
    """Try body sizes from MAX_SIZE down to MIN_SIZE and pick the one that needs the
    fewest pages (largest size wins a tie). Sizes where a line would be too wide for
    the page (and wrap, breaking the chord alignment) are skipped when possible.
    Returns (size, pages)."""
    prepared = prepare(sections, transpose, flats)
    meas = Measurer()
    text_w = (PAGE_WIDTH_CM - 2 * MARGIN_CM) * PT_PER_CM
    if autofit:
        n = int(round((MAX_SIZE - MIN_SIZE) / SIZE_STEP))
        sizes = [MAX_SIZE - i * SIZE_STEP for i in range(n + 1)]
    else:
        sizes = [size]
    results = []
    for sz in sizes:
        secs, widest = _pdf_sections(prepared, sz, meas, text_w)
        _, pages = _layout(secs, sz, True, _default_geom())
        results.append((widest <= text_w, pages, sz))
    fitting = [r for r in results if r[0]]
    best = min(fitting, key=lambda r: (r[1], -r[2])) if fitting else results[-1]
    return best[2], best[1]


def make_pdf(title, artist, sections, path, transpose=0, flats=None, size=BODY_SIZE):
    """Same layout as the Word file, drawn directly (no Word needed).
    Uses Helvetica, which has exactly Arial's letter widths."""
    text_w = (PAGE_WIDTH_CM - 2 * MARGIN_CM) * PT_PER_CM
    prepared = prepare(sections, transpose, flats)
    secs, _ = _pdf_sections(prepared, size, Measurer(), text_w)
    return _write_pdf(path, f"{title} - {artist}", f"{title} \u2013 {artist}", secs, body_size=size)


# --------------------------------------------------------------------------
# Transposing an existing .docx
# --------------------------------------------------------------------------
def _is_chord_par(par):
    rs = [r for r in par.runs if r.text.strip()]
    if not rs or not all(r.bold for r in rs):
        return False
    if any(r.font.size and r.font.size.pt >= TITLE_SIZE for r in rs):
        return False                          # the title
    tokens = " ".join(r.text for r in rs).split()
    return bool(tokens) and all(is_chord(t) for t in tokens)


def transpose_docx(src, steps, dst, flats=None):
    from docx import Document
    doc = Document(src)

    chord_pars = [p for p in doc.paragraphs if _is_chord_par(p)]
    if flats is None:
        toks = [t for p in chord_pars for t in p.text.split() if is_chord(t)]
        flats = auto_flats([transpose_chord(t, steps, False) for t in toks])

    pattern = re.compile(r"(\S+)([ ]*)")
    for par in chord_pars:
        for r in par.runs:
            if not r.text.strip():
                continue

            def repl(m):
                old, spaces = m.group(1), m.group(2)
                if not is_chord(old):
                    return m.group(0)
                new = transpose_chord(old, steps, flats)
                # keep the following chord in the same column in space-aligned docs
                if spaces:
                    spaces = " " * max(1, len(spaces) + len(old) - len(new))
                return new + spaces
            r.text = pattern.sub(repl, r.text)
    doc.save(dst)
    return len(chord_pars)


def docx_to_pdf(src, dst):
    """PDF of a chord-sheet .docx (made by this tool or by hand): reads the
    paragraphs, tab stops and bold runs and draws them with the same layout rules
    (sections kept together on a page)."""
    from docx import Document
    doc = Document(src)
    meas = Measurer()
    pars = doc.paragraphs

    def items(par):
        stops = sorted(ts.position.pt for ts in par.paragraph_format.tab_stops)
        x, out = 0.0, []
        for r in par.runs:
            bold = bool(r.bold)
            size = r.font.size.pt if r.font.size else BODY_SIZE
            for piece in re.split(r"(\t)", r.text):
                if piece == "\t":
                    nxt = [t for t in stops if t > x + 0.5]
                    x = nxt[0] if nxt else (int(x // 36) + 1) * 36
                elif piece:
                    out.append((x, piece, bold, size))
                    x += meas.width(piece, bold, size)
        return out

    title_text, start = None, 0
    if pars and pars[0].text.strip() and any(r.font.size and r.font.size.pt >= TITLE_SIZE for r in pars[0].runs):
        title_text, start = pars[0].text.strip(), 1

    sizes = [r.font.size.pt for p in pars[start:] for r in p.runs if r.text.strip() and r.font.size]
    body_size = max(set(sizes), key=sizes.count) if sizes else BODY_SIZE

    sections, units, i = [], [], start          # sections: [units, blank_rows_after]
    while i < len(pars):
        par = pars[i]
        if not par.text.strip():                 # blank paragraph = gap after the section
            if units:
                sections.append([units, 0])
                units = []
            if sections:
                sections[-1][1] += 1
            else:
                sections.append([[], 1])
            i += 1
            continue
        group = [items(par)]
        if _is_chord_par(par) and i + 1 < len(pars) and pars[i + 1].text.strip() \
                and not _is_chord_par(pars[i + 1]):
            group.append(items(pars[i + 1]))     # chord line stays with its lyric line
            i += 1
        units.append(group)
        i += 1
    if units:
        sections.append([units, 0])
    s0 = doc.sections[0]
    geom = (s0.page_width.pt, s0.page_height.pt, s0.left_margin.pt, s0.top_margin.pt, s0.bottom_margin.pt)
    return _write_pdf(dst, os.path.splitext(os.path.basename(dst))[0], title_text, sections, geom, body_size)


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------
def safe_filename(s):
    return re.sub(r'[\\/:*?"<>|]+', "", s).strip()


def cmd_make(a):
    meta_capo = 0
    if a.url:
        title, artist, content, meta_capo = parse_ug_html(fetch_html(a.url))
    elif a.html:
        with open(a.html, encoding="utf-8", errors="replace") as f:
            title, artist, content, meta_capo = parse_ug_html(f.read())
    elif a.text:
        with open(a.text, encoding="utf-8") as f:
            content = f.read()
        title, artist = a.title or "Unknown Song", a.artist or "Unknown Artist"
    else:
        sys.exit("Give a URL, or --html FILE, or --text FILE.")
    title, artist = a.title or title, a.artist or artist

    capo = detect_capo(content, meta_capo)
    shift = a.transpose + (0 if a.keep_capo else capo)
    if capo:
        print(f"Capo {capo} found: " + ("kept as on the page." if a.keep_capo else
              f"chords raised by {capo} semitones to the original key."))

    sections = build_sections(content)
    flats = True if a.flats else (False if a.sharps else None)
    size, pages = plan_layout(title, artist, sections, shift, flats,
                              autofit=a.size is None, size=a.size or BODY_SIZE)
    print(f"Text size {size:g} pt -> {pages} page{'s' if pages != 1 else ''}.")
    out = a.output or f"{safe_filename(title)} - {safe_filename(artist)}.docx"
    make_docx(title, artist, sections, out, transpose=shift, flats=flats, size=size)
    print(f"Saved: {out}")
    if a.pdf:
        pdf = os.path.splitext(out)[0] + ".pdf"
        make_pdf(title, artist, sections, pdf, transpose=shift, flats=flats, size=size)
        print(f"Saved: {pdf}")


def cmd_transpose(a):
    out = a.output or a.file
    n = transpose_docx(a.file, int(a.steps), out,
                       flats=True if a.flats else (False if a.sharps else None))
    print(f"Transposed {n} chord lines by {int(a.steps):+d} semitones -> {out}")
    if a.pdf:
        pdf = os.path.splitext(out)[0] + ".pdf"
        docx_to_pdf(out, pdf)
        print(f"Saved: {pdf}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    m = sub.add_parser("make", help="build a Word document from an Ultimate Guitar page")
    m.add_argument("url", nargs="?", help="Ultimate Guitar chords URL")
    m.add_argument("--html", help="a saved Ultimate Guitar page instead of a URL")
    m.add_argument("--text", help="plain-text chord sheet ([ch]..[/ch] tags optional)")
    m.add_argument("--title"), m.add_argument("--artist")
    m.add_argument("-o", "--output")
    m.add_argument("--transpose", type=int, default=0, metavar="N", help="semitones, e.g. -2 or 3")
    m.add_argument("--pdf", action="store_true", help="also save a PDF next to the Word file")
    m.add_argument("--size", type=float, metavar="PT",
                   help="fixed text size, e.g. 12 (default: auto-fit between 8 and 12 for the fewest pages)")
    m.add_argument("--keep-capo", action="store_true",
                   help="keep the capo/guitar chords instead of converting to the original key")
    g = m.add_mutually_exclusive_group()
    g.add_argument("--flats", action="store_true"), g.add_argument("--sharps", action="store_true")
    m.set_defaults(func=cmd_make)

    t = sub.add_parser("transpose", help="transpose the chords of an existing .docx")
    t.add_argument("file"), t.add_argument("steps", help="semitones, e.g. +2 or -3")
    t.add_argument("-o", "--output", help="default: overwrite the file")
    t.add_argument("--pdf", action="store_true", help="also save a PDF of the transposed document")
    g = t.add_mutually_exclusive_group()
    g.add_argument("--flats", action="store_true"), g.add_argument("--sharps", action="store_true")
    t.set_defaults(func=cmd_transpose)

    a = ap.parse_args()
    try:
        a.func(a)
    except RuntimeError as e:
        sys.exit(str(e))


if __name__ == "__main__":
    main()